"""Word files (.docx), drawn with python-docx: A4, the flat lockup, then the paper's blocks (P9).

Every text is set as a run's text, after `printable`: python-docx writes it into the XML as text,
escaped, never as markup. Word can't embed Newsreader and Geist, and few readers have them, so a
file uses Georgia and Arial, which every Word install has, in the same roles.
"""

from io import BytesIO

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph as WordParagraph

from aiact.export.paper import Block, FillMarker, Heading, Note, Paper, Paragraph, Quote
from aiact.export.style import (
    FILL_TINT,
    INK,
    INK_2,
    LENS_BLUE,
    LOGO_PNG,
    LOGO_WIDTH_MM,
    MARGIN_MM,
    QUOTE_TINT,
    SLATE,
    STAMP,
)
from aiact.export.text import printable

SERIF, SANS = "Georgia", "Arial"
A4_WIDTH_MM, A4_HEIGHT_MM = 210, 297


def render_docx(paper: Paper) -> bytes:
    """`paper` as a Word file."""
    document = Document()
    _set_up(document, paper)
    document.add_picture(str(LOGO_PNG), width=Mm(LOGO_WIDTH_MM))
    if paper.stamp is not None:
        _text(document.add_paragraph(), paper.stamp, SANS, 9, STAMP, bold=True)
    _text(document.add_paragraph(), paper.title, SERIF, 22, INK)
    for line in paper.lines:
        _text(document.add_paragraph(), line, SANS, 9, SLATE)
    for block in paper.blocks:
        _block(document, block)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _set_up(document: WordDocument, paper: Paper) -> None:
    """A4 with the paper's margins, its title and language, and the footer on every page."""
    section = document.sections[0]
    section.page_width, section.page_height = Mm(A4_WIDTH_MM), Mm(A4_HEIGHT_MM)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(MARGIN_MM))
    document.core_properties.title = printable(paper.title)
    document.core_properties.language = paper.language.value
    footer = section.footer.paragraphs[0]
    _text(footer, paper.footer, SANS, 8, SLATE)
    page = section.footer.add_paragraph()
    page.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    before, rest = paper.page_of.split("{page}", 1)
    between, after = rest.split("{pages}", 1)
    for text, field in ((before, "PAGE"), (between, "NUMPAGES"), (after, None)):
        _text(page, text, SANS, 8, SLATE)
        if field is not None:
            _field(page, field)


def _block(document: WordDocument, block: Block) -> None:
    paragraph = document.add_paragraph()
    match block:
        case Heading(text=text):
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.keep_with_next = True
            _text(paragraph, text, SANS, 11, INK, bold=True)
        case Paragraph(text=text):
            _text(paragraph, text, SERIF, 11.5, INK_2)
        case Quote(text=text, source=source):
            _shade(paragraph, QUOTE_TINT)
            _text(paragraph, text, SERIF, 10.5, INK_2)
            _text(document.add_paragraph(), source, SANS, 8, SLATE)
        case FillMarker(text=text):
            _shade(paragraph, FILL_TINT)
            _text(paragraph, text, SANS, 10, LENS_BLUE, bold=True)
        case Note(text=text):
            _shade(paragraph, QUOTE_TINT)
            _text(paragraph, text, SANS, 9.5, INK_2)


def _text(
    paragraph: WordParagraph,
    text: str,
    font: str,
    size: float,
    colour: str,
    *,
    bold: bool = False,
) -> None:
    """`text` as given, as a run of the paragraph; a line break in it stays a line break."""
    run = paragraph.add_run(printable(text))
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(colour.removeprefix("#"))


def _shade(paragraph: WordParagraph, colour: str) -> None:
    """A ground of `colour` behind the paragraph."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), colour.removeprefix("#"))
    paragraph._p.get_or_add_pPr().append(shading)  # pyright: ignore[reportPrivateUsage]


def _field(paragraph: WordParagraph, instruction: str) -> None:
    """A field Word fills in, such as the page number (`PAGE`) or the page count (`NUMPAGES`)."""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    paragraph._p.append(field)  # pyright: ignore[reportPrivateUsage]
