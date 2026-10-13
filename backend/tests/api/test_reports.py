"""`POST /reports/result` and `/reports/inventory` (P25): the report for a lawyer made from a
posted result, validated against the contract, and the inventory report, as PDF and Word files
made in the export workers."""

import logging
from collections.abc import Callable, Iterator
from datetime import date
from io import BytesIO
from typing import Any

import httpx
import pytest
from docx import Document
from fastapi import FastAPI

from aiact.api import create_app
from aiact.api.dependencies import ApiContext
from aiact.api.files import block_texts
from aiact.api.reports import (
    INVENTORY_MAX_CHARS,
    INVENTORY_MAX_LINES,
    RENDERERS,
    ChecklistRequest,
    InventoryReportRequest,
    report_date,
)
from aiact.export import Paper, pool, render_pdf
from aiact.export.pdf import KlarPdf
from aiact.export.pool import RenderTimeout, RenderWorker
from aiact.obligations.list import default_table, list_obligations
from aiact.report.common import MAX_NAME_CHARS
from aiact.schemas import (
    DISCLAIMER_EN,
    ClassificationResult,
    EvidenceTag,
    Fact,
    HighRiskRoute,
    Obligation,
    RiskClass,
    Role,
)
from tests.api.conftest import BASE_URL, api_settings
from tests.drafting.builders import limited_risk_result
from tests.export.reading import one_line, pdf_pages, read_docx
from tests.schemas.builders import needs_description, rebuild, stored_before_decision_030
from tests.schemas.worked_example import worked_example_result

pytestmark = pytest.mark.anyio

PREPARED_ON = date(2026, 10, 1)
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
APPLIES_IF = {"en": "Applies if:", "de": "Gilt, wenn:"}
"""What comes before a duty's condition in the report, by language."""
SECRET = "Confidential: the Hamburg pilot of Project Falcon"
"""A name the logs and errors must never show."""


@pytest.fixture(autouse=True)
def fixed_date(app: FastAPI) -> Iterator[None]:
    app.dependency_overrides[report_date] = lambda: PREPARED_ON
    yield


@pytest.fixture
def report_app() -> Iterator[Callable[..., FastAPI]]:
    """Apps with their own limits; their export workers are stopped after the test."""
    apps: list[FastAPI] = []

    def build(**changes: object) -> FastAPI:
        app = create_app(api_settings(**changes))
        app.dependency_overrides[report_date] = lambda: PREPARED_ON
        apps.append(app)
        return app

    yield build
    for app in apps:
        app.state.context.export_slots.close()


def _body(result: ClassificationResult | None = None, **fields: Any) -> dict[str, Any]:
    dumped = (result or worked_example_result()).model_dump(mode="json")
    return {"result": dumped, "language": "en", "format": "pdf", **fields}


AMENDED = "amended"
"""The checklist's law version when a test isn't about it."""


def _tool(**fields: Any) -> dict[str, Any]:
    return {
        "name": "Applicant screening",
        "risk_class": "high_risk",
        "status": "decided",
        "role": "deployer",
        "law_version": "amended",
        "as_of": "2026-09-27",
        "duties": [{"label": "Keep the logs", "article": "art:26:6", "applies_from": "2027-12-02"}],
        **fields,
    }


async def test_the_worked_example_is_a_two_page_pdf_of_the_posted_results_class(
    client: httpx.AsyncClient,
) -> None:
    body = _body(tool_name="Applicant screening", company="Nordwerk Logistik GmbH")
    response = await client.post("/reports/result", json=body)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    # A constant name: the browser saves it under its own (`resultReportFileName`).
    assert response.headers["content-disposition"] == 'attachment; filename="klar-report.pdf"'
    first, second = pdf_pages(response.content)
    assert "Nordwerk Logistik GmbH. Prepared on 1 October 2026." in first
    assert "Class Status Role Duties apply from High-risk Decided Deployer" in first
    assert "How the result was reached" in first
    assert "Duties of the deployer, from 2 December 2027" in second
    assert "Points for legal review" in second
    for page in (first, second):
        assert DISCLAIMER_EN in page


async def test_the_class_follows_the_posted_result(client: httpx.AsyncClient) -> None:
    response = await client.post("/reports/result", json=_body(limited_risk_result("art:50:1")))
    assert response.status_code == 200
    assert "Duties apply from Limited risk Decided" in pdf_pages(response.content)[0]


async def test_a_german_word_report(client: httpx.AsyncClient) -> None:
    response = await client.post("/reports/result", json=_body(language="de", format="docx"))
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX
    word = read_docx(response.content)
    assert "Pflichten des Betreibers, ab dem 2. Dezember 2027" in word.paragraphs
    assert word.language == "de"


async def test_a_result_with_no_class_is_a_422(client: httpx.AsyncClient) -> None:
    response = await client.post("/reports/result", json=_body(needs_description()))
    assert response.status_code == 422
    assert response.json() == {
        "detail": "This result has no class yet, so there is no report to make. "
        "Describe the tool first."
    }


async def test_a_result_with_a_check_not_assessed_is_a_422_that_says_so(
    client: httpx.AsyncClient,
) -> None:
    """A result stored before decision 030 lists a refused check next to the class it assumed:
    that class rests on a check nobody assessed, so no report states it."""
    stored = stored_before_decision_030(worked_example_result())
    response = await client.post("/reports/result", json=_body(stored))
    assert response.status_code == 422
    assert response.json() == {
        "detail": "A check of this result was not assessed, so it has no class to report. "
        "Classify the tool again, or have a person review that check."
    }


async def test_a_result_that_breaks_the_contract_is_refused_without_echoing_it(
    client: httpx.AsyncClient, caplog: pytest.LogCaptureFixture
) -> None:
    forged = {**_body()["result"], "risk_class": "prohibited", "role": SECRET}
    with caplog.at_level(logging.DEBUG):
        response = await client.post(
            "/reports/result", json=_body(result=None) | {"result": forged}
        )
    assert response.status_code == 422
    assert response.json() == {"detail": "The result is not a valid classification result."}
    assert SECRET not in caplog.text


async def test_names_are_capped(client: httpx.AsyncClient) -> None:
    long = await client.post("/reports/result", json=_body(tool_name="x" * 201))
    fits = await client.post("/reports/result", json=_body(tool_name="x" * 200))
    assert (long.status_code, fits.status_code) == (422, 200)


async def test_a_name_never_reaches_the_logs(
    client: httpx.AsyncClient, caplog: pytest.LogCaptureFixture
) -> None:
    with caplog.at_level(logging.DEBUG):
        response = await client.post(
            "/reports/result", json=_body(tool_name=SECRET, company=SECRET)
        )
    assert response.status_code == 200
    assert SECRET not in caplog.text


async def test_the_inventory_report(client: httpx.AsyncClient) -> None:
    body = {
        "tools": [_tool(), _tool(name=None, risk_class=None, status=None, role=None, duties=[])],
        "language": "en",
        "format": "pdf",
        "company": "Nordwerk Logistik GmbH",
    }
    response = await client.post("/reports/inventory", json=body)
    assert response.status_code == 200
    assert response.headers["content-disposition"] == (
        'attachment; filename="klar-inventory-report.pdf"'
    )
    text = " ".join(pdf_pages(response.content))
    assert "Applicant screening High-risk Decided Deployer 2 December 2027" in text
    assert "Your tool Not classified yet" in text
    # The table's label for the duty, not the posted one.
    assert "Keep the logs it creates for at least six months Art. 26(6) 2 December 2027" in text


async def test_the_inventory_reports_caps(client: httpx.AsyncClient) -> None:
    body: dict[str, Any] = {"language": "en", "format": "docx"}
    too_many = await client.post("/reports/inventory", json={**body, "tools": [_tool()] * 101})
    bad_article = await client.post(
        "/reports/inventory",
        json={
            **body,
            "tools": [
                _tool(duties=[{"label": "x", "article": "26", "applies_from": "2027-12-02"}])
            ],
        },
    )
    too_many_duties = await client.post(
        "/reports/inventory",
        json={**body, "tools": [_tool(duties=_tool()["duties"] * 101)]},
    )
    assert (too_many.status_code, bad_article.status_code, too_many_duties.status_code) == (
        422,
        422,
        422,
    )


@pytest.mark.parametrize(("length", "status_code"), [(1_000, 200), (1_001, 422)])
async def test_a_duty_label_is_capped_at_1000_characters(
    client: httpx.AsyncClient, length: int, status_code: int
) -> None:
    duty = {"label": "x" * length, "article": "art:26:6", "applies_from": "2027-12-02"}
    body = {"tools": [_tool(duties=[duty])], "language": "en", "format": "docx"}
    assert (await client.post("/reports/inventory", json=body)).status_code == status_code


def test_the_report_caps_are_pinned() -> None:
    assert (
        InventoryReportRequest.MAX_TOOLS,
        INVENTORY_MAX_LINES,
        INVENTORY_MAX_CHARS,
        MAX_NAME_CHARS,
    ) == (100, 11_500, 330_000, 200)


async def test_a_hundred_tools_of_a_dozen_duties_each_fit(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(RenderWorker, "run", lambda _worker, _fn, _paper: b"%PDF")
    duty = {
        "label": "Keep the logs it creates for at least six months",
        "article": "art:26:6",
        "applies_from": "2027-12-02",
    }
    tools = [_tool(name=f"Tool {number}", duties=[duty] * 12) for number in range(100)]
    body = {"tools": tools, "language": "de", "format": "pdf"}
    assert (await client.post("/reports/inventory", json=body)).status_code == 200


async def test_a_report_too_long_for_a_file_is_a_413_before_rendering(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    rendered: list[object] = []
    monkeypatch.setattr(
        RenderWorker, "run", lambda _worker, _fn, paper: rendered.append(paper) or b"%"
    )
    # 4,000 duties the table has (so each prints its own label), three cells each: over the
    # line cap, within the body cap.
    duty = {"label": "x", "article": "art:26:6", "applies_from": "2027-12-02"}
    body = {
        "tools": [_tool(duties=[duty] * 100)] * 40,
        "language": "en",
        "format": "pdf",
    }
    response = await client.post("/reports/inventory", json=body)
    assert response.status_code == 413
    assert response.json() == {"detail": "This report is too long to make into a file."}
    assert rendered == []


async def test_reports_share_the_exports_rate_limit(report_app: Callable[..., FastAPI]) -> None:
    app = report_app(rate_limit_exports=1)
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url=BASE_URL) as client:
        first = await client.post("/reports/result", json=_body(format="docx"))
        second = await client.post(
            "/reports/inventory", json={"tools": [], "language": "en", "format": "docx"}
        )
    assert (first.status_code, second.status_code) == (200, 429)


async def test_a_report_with_every_slot_taken_is_a_503_that_gives_back_its_count(
    report_app: Callable[..., FastAPI],
) -> None:
    app = report_app(rate_limit_exports=1, export_max_concurrent=1)
    context: ApiContext = app.state.context
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url=BASE_URL) as client:
        taken = context.export_slots.try_acquire()
        assert taken is not None
        busy = await client.post("/reports/result", json=_body(format="docx"))
        context.export_slots.release(taken)
        after = await client.post("/reports/result", json=_body(format="docx"))
    assert (busy.status_code, busy.headers["Retry-After"]) == (503, "5")
    assert after.status_code == 200


async def test_a_report_past_its_time_limit_is_a_504(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    def timed_out(_worker: RenderWorker, *_args: object) -> bytes:
        raise RenderTimeout

    monkeypatch.setattr(RenderWorker, "run", timed_out)
    response = await client.post("/reports/result", json=_body())
    assert (response.status_code, response.json()) == (
        504,
        {"detail": "Making this report took too long. Please try again."},
    )


@pytest.mark.parametrize("point", ["annex:III", "x", "annex:I:", "annex:III:99:z"])
async def test_an_annex_point_the_report_cant_name_still_makes_a_report(
    client: httpx.AsyncClient, point: str
) -> None:
    result = worked_example_result().model_copy(update={"annex_point": point})
    response = await client.post("/reports/result", json=_body(result))
    assert response.status_code == 200


def _with_tall_fact(lines: int) -> ClassificationResult:
    tall = Fact(text=f"{SECRET}\n" * (lines - 1) + SECRET, evidence=EvidenceTag.ASSUMED)
    worked = worked_example_result()
    return worked.model_copy(update={"deciding_facts": (tall, *worked.deciding_facts)})


async def test_a_cell_taller_than_a_page_is_a_413_that_keeps_its_count(
    report_app: Callable[..., FastAPI], caplog: pytest.LogCaptureFixture
) -> None:
    """fpdf2 can't break a table row across pages: the renderer's own error is a 413 saying the
    report is too long, and the count stays (the render took a worker and its time)."""
    app = report_app(rate_limit_exports=1)
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url=BASE_URL) as client:
        with caplog.at_level(logging.DEBUG):
            refused = await client.post("/reports/result", json=_body(_with_tall_fact(80)))
        after = await client.post("/reports/result", json=_body(format="docx"))
    assert (refused.status_code, refused.json()) == (
        413,
        {"detail": "This report is too long to make into a file."},
    )
    assert after.status_code == 429
    assert SECRET not in caplog.text


async def test_a_renderer_error_is_a_500_logged_as_an_error_by_its_type_only(
    report_app: Callable[..., FastAPI],
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """Klar's fault, not the text's: not "too long", and logged at error level, with no text of
    the report. The count stays, as for a 413: the render took a worker and its time."""
    monkeypatch.setitem(RENDERERS, "pdf", int)  # int(paper) raises in the worker process
    app = report_app(rate_limit_exports=1)
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url=BASE_URL) as client:
        with caplog.at_level(logging.DEBUG):
            failed = await client.post(
                "/reports/result", json=_body(tool_name=SECRET, company=SECRET)
            )
        after = await client.post("/reports/result", json=_body(format="docx"))
    assert (failed.status_code, failed.json()) == (
        500,
        {"detail": "Klar couldn't make this file. Please try again later."},
    )
    assert after.status_code == 429
    errors = [r.getMessage() for r in caplog.records if r.levelno == logging.ERROR]
    assert errors == ["export failed: the render raised TypeError"]
    assert SECRET not in caplog.text


def render_pdf_with_another_table_error(paper: Paper) -> bytes:
    """`render_pdf` with its table draw raising a `ValueError` that isn't fpdf2's "row too
    high". Run in a worker process, which a test's monkeypatch doesn't reach, so it patches
    there."""

    def broken(*_args: object, **_kwargs: object) -> None:
        raise ValueError("something else")

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(KlarPdf, "_table", broken)
        return render_pdf(paper)


async def test_another_table_error_is_a_500_not_a_413(
    report_app: Callable[..., FastAPI],
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    monkeypatch.setitem(RENDERERS, "pdf", render_pdf_with_another_table_error)
    app = report_app()
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url=BASE_URL) as client:
        with caplog.at_level(logging.DEBUG):
            failed = await client.post("/reports/result", json=_body())
    assert failed.status_code == 500
    errors = [r.getMessage() for r in caplog.records if r.levelno == logging.ERROR]
    assert errors == ["export failed: the render raised ValueError"]


async def test_a_worker_that_doesnt_start_is_a_503_that_gives_back_its_count(
    report_app: Callable[..., FastAPI], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(pool, "START_TIMEOUT_SECONDS", 0.0)  # no new worker is ready at once
    app = report_app(rate_limit_exports=1)
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url=BASE_URL) as client:
        busy = await client.post("/reports/result", json=_body(format="docx"))
        monkeypatch.undo()
        after = await client.post("/reports/result", json=_body(format="docx"))
    assert (busy.status_code, busy.headers["Retry-After"]) == (503, "5")
    assert after.status_code == 200, "the count was given back"


def _every_table_duty() -> tuple[ClassificationResult, ...]:
    """For each role, a high-risk and a limited-risk result listing every duty the committed
    table gives that role (on every route): together, every entry a result report can print."""
    high, limited = worked_example_result(), limited_risk_result("art:50:1")
    return tuple(
        rebuild(base, role=role, obligations=list_obligations(role, risk_class).obligations)
        for role in Role
        for base, risk_class in ((high, RiskClass.HIGH_RISK), (limited, RiskClass.LIMITED_RISK))
    )


def test_every_reportable_table_entry_is_posted() -> None:
    posted = {(duty.article, duty.duty) for r in _every_table_duty() for duty in r.obligations}
    reportable = {
        (entry.article, entry.duty)
        for entry in default_table().obligations
        if {RiskClass.HIGH_RISK, RiskClass.LIMITED_RISK} & set(entry.risk_classes)
    }
    assert reportable <= posted
    assert (
        "art:27:1",
        next(e.duty for e in default_table().obligations if e.article == "art:27:1"),
    ) in posted


def _printed_duty(duty: Obligation, language: str) -> str:
    """A posted table duty as the report prints it: its formal wording in `language`, then its
    condition in `language` too, on one line."""
    text = duty.duty_de if language == "de" else duty.duty
    if duty.applies_if is not None:
        condition = duty.applies_if_de if language == "de" else duty.applies_if
        text = f"{text} {APPLIES_IF[language]} {condition}"
    return one_line(text)


@pytest.mark.parametrize("language", ["en", "de"])
async def test_every_table_duty_makes_a_report(
    report_app: Callable[..., FastAPI], language: str
) -> None:
    """The long duties with their conditions (Art. 27(1), Art. 26(10), Art. 53(1)(b)) fit a
    page, and each one reaches the file: no deployer of a credit score or a public body is
    refused its report or loses a duty from it."""
    app = report_app(rate_limit_exports=100)
    transport = httpx.ASGITransport(app=app)
    printed: set[str] = set()
    async with httpx.AsyncClient(transport=transport, base_url=BASE_URL) as client:
        for result in _every_table_duty():
            response = await client.post("/reports/result", json=_body(result, language=language))
            assert response.status_code == 200, (result.role, result.risk_class)
            text = " ".join(pdf_pages(response.content))
            for duty in result.obligations:
                assert _printed_duty(duty, language) in text, (result.role, duty.article)
                printed.add(duty.article)
    fria = next(e for e in default_table().obligations if e.article == "art:27:1")
    assert fria.applies_if is not None
    assert "art:27:1" in printed


def test_the_report_date_is_today() -> None:
    assert report_date() == date.today()


def _checklist_tool() -> dict[str, Any]:
    """Applicant screening with one duty ticked off and one open."""
    return _tool(
        duties=[
            {"label": "Keep the logs", "article": "art:26:6", "applies_from": "2027-12-02"}
            | {"is_done": True},
            {"label": "Use it as instructed", "article": "art:26:1", "applies_from": "2027-12-02"},
        ]
    )


async def test_the_checklist_is_a_pdf_with_each_duties_ticked_state(
    client: httpx.AsyncClient,
) -> None:
    body = {"tools": [_checklist_tool()], "language": "en", "format": "pdf", "law_version": AMENDED}
    response = await client.post("/reports/checklist", json=body)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == (
        'attachment; filename="klar-obligations-checklist.pdf"'
    )
    text = " ".join(pdf_pages(response.content))
    assert "Applicant screening High-risk, Deployer. 1 of 2 done." in text
    assert "Done Keep the logs it creates for at least six months Art. 26(6)" in text
    assert "Open Use the tool as the vendor's instructions say Art. 26(1)" in text


async def test_the_checklist_is_a_word_file_with_each_duties_ticked_state(
    client: httpx.AsyncClient,
) -> None:
    body = {
        "tools": [_checklist_tool()],
        "language": "de",
        "format": "docx",
        "law_version": AMENDED,
    }
    response = await client.post("/reports/checklist", json=body)
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX
    assert response.headers["content-disposition"] == (
        'attachment; filename="klar-obligations-checklist.docx"'
    )
    [table] = Document(BytesIO(response.content)).tables
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    assert rows[0] == ["Stand", "Pflicht", "Vorschrift", "Gilt ab"]
    assert [(row[0], row[2]) for row in rows[1:]] == [
        ("Erledigt", "Art. 26 Abs. 6"),
        ("Offen", "Art. 26 Abs. 1"),
    ]
    assert "1 von 2 erledigt." in " ".join(read_docx(response.content).paragraphs)


async def test_the_checklist_keeps_the_inventory_reports_caps(client: httpx.AsyncClient) -> None:
    body: dict[str, Any] = {"language": "en", "format": "docx", "law_version": AMENDED}
    too_many = await client.post(
        "/reports/checklist", json={**body, "tools": [_checklist_tool()] * 101}
    )
    not_a_flag = await client.post(
        "/reports/checklist",
        json={**body, "tools": [_tool(duties=[_tool()["duties"][0] | {"is_done": "maybe"}])]},
    )
    other_law = await client.post(
        "/reports/checklist", json={**body, "tools": [_checklist_tool()], "law_version": "draft"}
    )
    no_law = await client.post(
        "/reports/checklist",
        json={"language": "en", "format": "docx", "tools": [_checklist_tool()]},
    )
    assert (too_many.status_code, not_a_flag.status_code) == (422, 422)
    assert (other_law.status_code, no_law.status_code) == (422, 422)
    for refused in (other_law, no_law):
        assert any(error["loc"][-1] == "law_version" for error in refused.json()["errors"])
    assert ChecklistRequest.model_fields["tools"].metadata[0].max_length == 100


async def test_the_checklist_under_the_original_text_takes_its_route_dates(
    client: httpx.AsyncClient,
) -> None:
    """The screen under ?law=original dates Annex III duties 2 August 2026 (P10's `routeDates`);
    the checklist downloaded from it does too, with the screen's notice and the law it follows.
    The amended export keeps the stored 2 December 2027 and has no notice."""
    tools = [_checklist_tool()]
    original = await client.post(
        "/reports/checklist",
        json={"tools": tools, "language": "en", "format": "pdf", "law_version": "original"},
    )
    amended = await client.post(
        "/reports/checklist",
        json={"tools": tools, "language": "en", "format": "pdf", "law_version": AMENDED},
    )
    original_text = " ".join(pdf_pages(original.content))
    amended_text = " ".join(pdf_pages(amended.content))
    notice = "Dates follow the original Art. 113. The duties shown are the amended text's."
    assert "Art. 26(6) 2 August 2026" in original_text
    assert "2 December 2027" not in original_text
    assert notice in original_text
    assert (
        "Dates under Regulation (EU) 2024/1689 as published on 12 July 2024 (CELEX 32024R1689)."
        in original_text
    )
    assert "Art. 26(6) 2 December 2027" in amended_text
    assert notice not in amended_text
    assert (
        "Dates under Regulation (EU) 2024/1689 as amended, consolidated text of 27 July 2026 "
        "(CELEX 02024R1689-20260727)."
    ) in amended_text


async def test_the_checklist_and_the_inventory_report_name_the_law_version(
    client: httpx.AsyncClient,
) -> None:
    """Each tool's block says which text its result was checked under, as of when."""
    checked = "Checked under the amended text, as it applies on 27 September 2026."
    checklist = await client.post(
        "/reports/checklist",
        json={"tools": [_checklist_tool()], "language": "en", "format": "pdf"}
        | {"law_version": AMENDED},
    )
    inventory = await client.post(
        "/reports/inventory", json={"tools": [_tool()], "language": "en", "format": "pdf"}
    )
    assert checked in " ".join(pdf_pages(checklist.content))
    assert checked in " ".join(pdf_pages(inventory.content))


# --- Relief notes for an SME or a small mid-cap (P31, user decision D6) ---------------------------


def _provider_tool() -> dict[str, Any]:
    duty = {
        "label": "Set up a quality management system",
        "article": "art:16:c",
        "applies_from": "2027-12-02",
    }
    return _tool(role="provider", duties=[duty])


@pytest.mark.parametrize(
    ("size_class", "is_shown"), [("sme", True), ("smc", True), ("none", False), (None, False)]
)
async def test_each_report_passes_the_size_class_on(
    client: httpx.AsyncClient, size_class: str | None, is_shown: bool
) -> None:
    """The profile's size class reaches every report: an SME's or a small mid-cap's lists the
    table's relief note under the duty."""
    duties = list_obligations(Role.PROVIDER, RiskClass.HIGH_RISK, HighRiskRoute.ANNEX_III)
    result = rebuild(worked_example_result(), role=Role.PROVIDER, obligations=duties.obligations)
    tools = {"tools": [_provider_tool()], "language": "en", "format": "pdf"}
    requests = {
        "/reports/result": _body(result, size_class=size_class),
        "/reports/inventory": {**tools, "size_class": size_class},
        "/reports/checklist": {**tools, "law_version": AMENDED, "size_class": size_class},
    }
    for path, body in requests.items():
        response = await client.post(path, json=body)
        assert response.status_code == 200, path
        text = one_line(" ".join(pdf_pages(response.content)))
        assert ("Relief for your company" in text) is is_shown, path


async def test_an_invalid_size_class_is_a_422_on_every_report(client: httpx.AsyncClient) -> None:
    tools = {"tools": [_provider_tool()], "language": "en", "format": "pdf"}
    requests = {
        "/reports/result": _body(size_class="large"),
        "/reports/inventory": {**tools, "size_class": "large"},
        "/reports/checklist": {**tools, "law_version": AMENDED, "size_class": "large"},
    }
    for path, body in requests.items():
        response = await client.post(path, json=body)
        assert response.status_code == 422, path
        assert response.json()["errors"][0]["loc"] == ["body", "size_class"], path


@pytest.mark.parametrize(
    ("path", "fields"),
    [("/reports/inventory", {}), ("/reports/checklist", {"law_version": AMENDED})],
)
async def test_a_hundred_high_risk_providers_of_an_sme_in_german_fit_the_file_caps(
    report_app: Callable[..., FastAPI],
    monkeypatch: pytest.MonkeyPatch,
    path: str,
    fields: dict[str, str],
) -> None:
    """The longest single-role duty list (26 duties), with the table's German labels and an
    SME's relief notes, for `MAX_TOOLS` tools, is made, not refused as too long. The default
    body cap admits about 80 such tools, so this app's is raised."""
    rendered: list[Paper] = []
    monkeypatch.setattr(
        RenderWorker, "run", lambda _worker, _fn, paper: rendered.append(paper) or b"%PDF"
    )
    duties = list_obligations(Role.PROVIDER, RiskClass.HIGH_RISK, HighRiskRoute.ANNEX_III)
    labelled = [
        {"label": duty.label_de, "article": duty.article, "applies_from": str(duty.applies_from)}
        for duty in duties.obligations
    ]
    assert len(labelled) == 26
    tools = [
        _tool(name=f"Werkzeug {number}", role="provider", duties=labelled) for number in range(100)
    ]
    body = {"tools": tools, "language": "de", "format": "pdf", "size_class": "sme", **fields}
    app = report_app(api_max_body_bytes=1_048_576)
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url=BASE_URL) as client:
        response = await client.post(path, json=body)
    assert response.status_code == 200
    (paper,) = rendered
    texts = [text for block in paper.blocks for text in block_texts(block)]
    # Two relief notes a tool: the quality management system and Annex IV.
    assert sum("Erleichterung für Ihr Unternehmen" in text for text in texts) == 200
