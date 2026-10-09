"""Word files (P9): every block as plain text, the logo, and the footer with the page fields."""

from io import BytesIO
from zipfile import ZipFile

from docx import Document

from aiact.export import render_docx
from aiact.schemas import Language
from tests.export.papers import FOOTER, LONG, MARKUP, report_paper, sample_paper
from tests.export.reading import read_docx

SIGNATURE_LINE = "_" * 48


def test_every_block_is_written_as_given_in_order() -> None:
    word = read_docx(render_docx(sample_paper()))
    texts = [text for text in word.paragraphs if text]
    assert texts == [
        "DRAFT",
        "Größere Dokumentation: a test",
        "AI Act text: CELEX 02024R1689-20260727, as it applies on 7 October 2026.",
        "A simplified form is expected.",
        "The system",
        MARKUP,
        "Line one\nline two",
        "the methods and steps performed for the development",
        "Annex IV 2",
        "To fill in.",
        LONG,
    ]


def test_markup_is_escaped_into_the_file_never_read_as_markup() -> None:
    with ZipFile(BytesIO(render_docx(sample_paper()))) as archive:
        body = archive.read("word/document.xml").decode()
    assert "&lt;b&gt;Not bold&lt;/b&gt; **not starred** {nb} &amp;amp;" in body
    assert "<script>" not in body


def test_the_footer_holds_the_notice_and_the_page_fields() -> None:
    word = read_docx(render_docx(sample_paper()))
    assert word.footer[0] == FOOTER
    assert word.footer[1] == "Page  of "
    assert 'w:instr="PAGE"' in word.footer_xml
    assert 'w:instr="NUMPAGES"' in word.footer_xml


def test_a_german_paper_is_marked_german_and_numbers_its_pages_in_german() -> None:
    word = read_docx(render_docx(sample_paper(Language.DE)))
    assert (word.language, word.footer[1]) == ("de", "Seite  von ")
    assert word.title == "Größere Dokumentation: a test"


def test_the_flat_lockup_is_the_only_picture() -> None:
    with ZipFile(BytesIO(render_docx(sample_paper()))) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert len(media) == 1
    assert media[0].endswith(".png")


def test_the_reports_blocks_are_word_tables_a_page_break_and_signature_lines() -> None:
    data = render_docx(report_paper())
    texts = [text for text in read_docx(data).paragraphs if text]
    assert texts == [
        "Classification report · Report KR-0000ABCD",
        "Applicant screening",
        "Prepared on 1 October 2026.",
        "Facts",
        "Duties",
        "For the reviewer",
        SIGNATURE_LINE,
        "Reviewed by",
        SIGNATURE_LINE,
        "Date",
        "Not legal advice.",
    ]
    tables = Document(BytesIO(data)).tables
    assert [[cell.text for cell in row.cells] for row in tables[0].rows] == [
        ["Class", "Status"],
        ["High-risk", "Decided"],
    ]
    assert [[cell.text for cell in row.cells] for row in tables[1].rows] == [
        ["Fact", "Source"],
        [MARKUP, "Stated by the user"],
    ]
    with ZipFile(BytesIO(data)) as archive:
        body = archive.read("word/document.xml").decode()
    assert body.count("<w:tblHeader/>") == 2, "each table's header repeats on a new page"
    assert '<w:br w:type="page"/>' in body
